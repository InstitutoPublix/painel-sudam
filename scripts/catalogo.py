"""Gera o catálogo dos indicadores: o que já está no painel, o que pode entrar e o que depende de decisão.

Uso:
    python scripts/catalogo.py                 # usa o .xlsx mais recente de data/
    python scripts/catalogo.py caminho.xlsx

Grava catalogo/catalogo_indicadores.xlsx (pasta fora do git). As colunas "automáticas" são recalculadas a cada
execução a partir da planilha e de config/indicadores.toml. As colunas "manuais" (Situação, Pendência, Decisão,
Responsável, Observações) são preenchidas à mão na própria planilha e preservadas ao gerar de novo, por ID.
"""
from __future__ import annotations

import argparse
import sys
from datetime import date
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_data as bd  # noqa: E402

OUT = bd.ROOT / "catalogo" / "catalogo_indicadores.xlsx"
MIN_COVER = bd.MIN_COVER
MANUAL = ["Situação (manual)", "Pendência / particularidade", "Decisão", "Responsável", "Observações", "Atualizado em"]
SITUACOES = ["No painel", "Pronto para entrar", "Aguardando decisão", "Ajuste no build", "Aguardando dado", "Descartado"]
AUTO_ORDER = ["ID", "Indicador", "Eixo", "Unidade", "Status da coleta", "No painel", "Diagnóstico automático",
              "Dado municipal", "Dado estadual", "Anos (município)", "Cobertura AL (último ano, %)",
              "Regra da ficha", "Regra suportada", "Ponderação (ficha)", "Fonte"]


def diagnose(row: dict) -> str:
    """Resumo dos motivos que impedem (ou não) a entrada automática no painel."""
    if row["Status da coleta"] != "ok":
        return f"Coleta não concluída: {str(row['Status da coleta'])[:70]}"
    if row["Dado municipal"] == "sem coluna" and row["Dado estadual"] in ("sem coluna", "vazio"):
        return "Sem coluna nas abas de recorte"
    if row["Dado municipal"] in ("sem coluna", "vazio"):
        return ("Só estadual/regional: entra, sem mapa municipal nem Amazônia Legal"
                if row["Dado estadual"] == "sim" else "Sem dado")
    issues = []
    if row["Cobertura AL (último ano, %)"] < MIN_COVER:
        issues.append(f"cobertura municipal baixa ({row['Cobertura AL (último ano, %)']:.0f}%)")
    if row["N anos"] < 2:
        issues.append("um único ano (sem série temporal)")
    if row["Regra suportada"] == "não":
        issues.append(f"regra de agregação sem suporte no build: '{row['Regra da ficha'][:40]}'")
    return "Candidato" if not issues else "; ".join(issues)


def collect(xlsx: Path) -> pd.DataFrame:
    book = pd.ExcelFile(xlsx, engine="openpyxl")
    lista = bd.read_lista(book)
    cfg = {i["id"] for i in bd.load_config()}
    frames = {lvl: book.parse(sheet) for lvl, sheet in bd.SHEETS.items()}
    cols = {lvl: {bd.norm(c): c for c in df.columns[df.columns.get_loc("Ano") + 1:]} for lvl, df in frames.items()}
    mun = frames["Município"]
    al = mun[mun["Amazônia Legal"].astype(str).str.strip().str.upper() == "SIM"]
    n_al = al.iloc[:, 0].nunique()
    col_cfg = {i["id"]: i.get("coluna") for i in bd.load_config()}

    rows = []
    for iid, f in lista.iterrows():
        name = col_cfg.get(iid) or f["Indicador"]
        key = bd.norm(name)

        def values(level):
            c = cols[level].get(key)
            return None if c is None else pd.to_numeric(frames[level][c], errors="coerce")

        m, e = values("Município"), values("Estado")
        rule = bd.agg_rule(str(f.get("Regra de Agregação Territorial", "")))
        r = {
            "ID": iid, "Indicador": f["Indicador"], "Eixo": bd.cell(f.get("Eixo P1")),
            "Unidade": bd.cell(f.get("Unidade de Medida")), "Status da coleta": bd.cell(f.get("Status")),
            "No painel": "sim" if iid in cfg else "",
            "Dado municipal": "sem coluna" if m is None else ("sim" if m.notna().any() else "vazio"),
            "Dado estadual": "sem coluna" if e is None else ("sim" if e.notna().any() else "vazio"),
            "Anos (município)": "", "Cobertura AL (último ano, %)": 0.0, "N anos": 0,
            "Regra da ficha": bd.cell(f.get("Regra de Agregação Territorial")),
            "Regra suportada": "sim" if rule else "não", "Ponderação (ficha)": bd.cell(f.get("Variável de Ponderação")),
            "Fonte": bd.cell(f.get("Fonte dos Dados"))[:80],
        }
        if m is not None and m.notna().any():
            sel = m.loc[al.index].notna()
            years = sorted(al.loc[sel, "Ano"].unique().astype(int))
            if years:
                r["N anos"] = len(years)
                r["Anos (município)"] = f"{years[0]}–{years[-1]}" if len(years) > 1 else str(years[0])
                last = al[al["Ano"] == years[-1]]
                r["Cobertura AL (último ano, %)"] = round(100 * last.loc[sel[last.index], last.columns[0]].nunique() / n_al, 1)
        r["Diagnóstico automático"] = diagnose(r)
        rows.append(r)
    return pd.DataFrame(rows)


def merge_manual(df: pd.DataFrame) -> pd.DataFrame:
    """Recupera as colunas manuais do catálogo anterior (por ID) e dá uma situação inicial quando em branco."""
    old = pd.read_excel(OUT, dtype=str).set_index("ID") if OUT.exists() else pd.DataFrame()
    for c in MANUAL:
        df[c] = df["ID"].map(old[c]) if c in old else None
    df["Situação (manual)"] = df["Situação (manual)"].where(df["Situação (manual)"].notna(), None)
    auto = df["No painel"].eq("sim").map({True: "No painel", False: None})
    df["Situação (manual)"] = df["Situação (manual)"].fillna(auto)
    # quem entrou no painel passa a constar como tal, mesmo que a situação anterior fosse outra
    df.loc[df["No painel"].eq("sim"), "Situação (manual)"] = "No painel"
    return df


def write(df: pd.DataFrame):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    OUT.parent.mkdir(exist_ok=True)
    order = AUTO_ORDER[:6] + ["Situação (manual)", "Pendência / particularidade", "Decisão", "Responsável",
                              "Observações", "Atualizado em"] + AUTO_ORDER[6:]
    df = df[order]
    with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Indicadores", index=False)
        resumo = (df.assign(**{"Situação (manual)": df["Situação (manual)"].fillna("(sem situação)")})
                  .groupby("Situação (manual)").size().rename("Indicadores").reset_index())
        resumo.to_excel(xw, sheet_name="Resumo", index=False)
        ws = xw.sheets["Indicadores"]
        manual = PatternFill("solid", fgColor="FFF2CC")
        head = PatternFill("solid", fgColor="1F3864")
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = manual if c.value in MANUAL else head
            if c.value in MANUAL:
                c.font = Font(bold=True, color="000000")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        widths = {"ID": 10, "Indicador": 48, "Eixo": 18, "Unidade": 22, "Status da coleta": 16, "No painel": 9,
                  "Situação (manual)": 20, "Pendência / particularidade": 50, "Decisão": 30, "Responsável": 14,
                  "Observações": 40, "Atualizado em": 13, "Diagnóstico automático": 50}
        for i, name in enumerate(order, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(name, 16)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        col = order.index("Situação (manual)") + 1
        dv = DataValidation(type="list", formula1='"' + ",".join(SITUACOES) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{ws.cell(2, col).column_letter}2:{ws.cell(len(df) + 1, col).column_letter}{len(df) + 1}")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("planilha", nargs="?", help="caminho do .xlsx (padrão: o mais recente em data/)")
    xlsx = bd.find_xlsx(ap.parse_args().planilha)
    print(f"Planilha: {xlsx.name}\nLendo abas (pode levar alguns minutos)…")
    df = merge_manual(collect(xlsx))
    write(df.drop(columns=["N anos"]))
    print(f"\nCatálogo gravado em {OUT} ({len(df)} indicadores, {date.today():%d/%m/%Y}).")
    print(df["Situação (manual)"].fillna("(sem situação)").value_counts().to_string())


if __name__ == "__main__":
    main()
